"""Single-pass: fill template + insert Word tables + format sub/superscripts.
No multiple file rewrites."""
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy
import re

doc = Document("/home/user/aegis/spinoff/Invention Disclosure Document.docx")
body = doc.element.body

with open("/home/user/aegis/spinoff/IDF_geometric_dosimetry.md") as f:
    md = f.read()

def extract_section(md, start_marker, end_marker):
    start = md.find(start_marker)
    if start == -1:
        return ""
    start = md.find('\n', start) + 1
    end = md.find(end_marker, start)
    if end == -1:
        end = len(md)
    return md[start:end].strip()

def get_full_text(elem):
    return ''.join(t.text for t in elem.iter(qn('w:t')) if t.text)

def make_run(parent, text, rPr_source=None, superscript=False, subscript=False, italic=False):
    r = parent.makeelement(qn('w:r'), {})
    if rPr_source is not None:
        rPr = deepcopy(rPr_source)
    else:
        rPr = r.makeelement(qn('w:rPr'), {})
    if superscript:
        va = rPr.makeelement(qn('w:vertAlign'), {})
        va.set(qn('w:val'), 'superscript')
        rPr.append(va)
    if subscript:
        va = rPr.makeelement(qn('w:vertAlign'), {})
        va.set(qn('w:val'), 'subscript')
        rPr.append(va)
    if italic:
        i_elem = rPr.find(qn('w:i'))
        if i_elem is None:
            rPr.append(rPr.makeelement(qn('w:i'), {}))
    r.append(rPr)
    t = r.makeelement(qn('w:t'), {})
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    return r

def process_formatted_text(text, parent, base_rPr):
    """Process text segment, creating runs with superscripts/subscripts/italics."""
    runs = []
    i = 0
    while i < len(text):
        # Superscript: digits^digits or letter^letters
        m = re.match(r'^(\d+)\^(\d+)', text[i:])
        if not m:
            m = re.match(r'^([a-zA-Z])\^([A-Za-z*]+)', text[i:])
        if m:
            if m.group(1):
                runs.append(make_run(parent, m.group(1), base_rPr))
            runs.append(make_run(parent, m.group(2), base_rPr, superscript=True))
            i += m.end()
            continue
        # Subscript: letter_word (not in middle of regular word, not filenames)
        m = re.match(r'^([A-Za-z])_([a-zA-Z]\w*)', text[i:])
        if m and (i == 0 or not text[i-1].isalnum()) and not re.match(r'.*\.(png|jpg|svg|txt)', text[i:i+40]):
            runs.append(make_run(parent, m.group(1), base_rPr))
            runs.append(make_run(parent, m.group(2), base_rPr, subscript=True))
            i += m.end()
            continue
        # Multi-letter subscript: SAR_wb
        m = re.match(r'^(SAR)_(\w+)', text[i:])
        if m:
            runs.append(make_run(parent, m.group(1), base_rPr))
            runs.append(make_run(parent, m.group(2), base_rPr, subscript=True))
            i += m.end()
            continue
        # Markdown italic
        m = re.match(r'^\*([^*]+)\*', text[i:])
        if m:
            runs.append(make_run(parent, m.group(1), base_rPr, italic=True))
            i += m.end()
            continue
        # >= <=
        if text[i:i+2] == '>=':
            runs.append(make_run(parent, '\u2265', base_rPr))
            i += 2
            continue
        if text[i:i+2] == '<=':
            runs.append(make_run(parent, '\u2264', base_rPr))
            i += 2
            continue
        # ~ before number
        if text[i] == '~' and i + 1 < len(text) and text[i+1].isdigit():
            runs.append(make_run(parent, '\u2248', base_rPr))
            i += 1
            continue
        # -- to en dash
        if text[i:i+4] == ' -- ':
            runs.append(make_run(parent, ' \u2013 ', base_rPr))
            i += 4
            continue
        # ||^2
        if text[i:i+4] == '||^2':
            runs.append(make_run(parent, '||', base_rPr))
            runs.append(make_run(parent, '2', base_rPr, superscript=True))
            i += 4
            continue
        # m^2
        if text[i:i+3] == 'm^2' and (i == 0 or text[i-1] == '/'):
            runs.append(make_run(parent, 'm', base_rPr))
            runs.append(make_run(parent, '2', base_rPr, superscript=True))
            i += 3
            continue
        # Regular text - accumulate
        j = i + 1
        while j < len(text):
            remaining = text[j:]
            if (re.match(r'^(\d+)\^(\d+)', remaining) or
                re.match(r'^([a-zA-Z])\^([A-Za-z*]+)', remaining) or
                re.match(r'^(SAR)_(\w+)', remaining) or
                (re.match(r'^([A-Za-z])_([a-zA-Z])', remaining) and (j == 0 or not text[j-1].isalnum()) and not re.match(r'.*\.(png|jpg|svg|txt)', remaining[:40])) or
                re.match(r'^\*[^*]+\*', remaining) or
                remaining[:2] in ('>=', '<=') or
                (remaining[0] == '~' and len(remaining) > 1 and remaining[1].isdigit()) or
                remaining[:4] == ' -- ' or
                remaining[:4] == '||^2' or
                (remaining[:3] == 'm^2' and text[j-1:j] == '/')):
                break
            j += 1
        runs.append(make_run(parent, text[i:j], base_rPr))
        i = j
    return runs

def replace_para_formatted(para_elem, new_text):
    """Replace paragraph text with formatted runs (sub/superscripts, italics)."""
    runs = para_elem.findall(qn('w:r'))
    base_rPr = None
    if runs:
        rPr = runs[0].find(qn('w:rPr'))
        if rPr is not None:
            base_rPr = deepcopy(rPr)
    for r in runs:
        para_elem.remove(r)
    for sdt in para_elem.findall(qn('w:sdt')):
        para_elem.remove(sdt)

    lines = new_text.split('\n')
    for li, line in enumerate(lines):
        if li > 0:
            br_run = para_elem.makeelement(qn('w:r'), {})
            if base_rPr is not None:
                br_run.append(deepcopy(base_rPr))
            br = br_run.makeelement(qn('w:br'), {})
            br_run.append(br)
            para_elem.append(br_run)
        for run in process_formatted_text(line, para_elem, base_rPr):
            para_elem.append(run)

def replace_para_plain(para_elem, new_text):
    """Replace paragraph text without formatting (for simple fields)."""
    runs = para_elem.findall(qn('w:r'))
    base_rPr = None
    if runs:
        rPr = runs[0].find(qn('w:rPr'))
        if rPr is not None:
            base_rPr = deepcopy(rPr)
    for r in runs:
        para_elem.remove(r)
    for sdt in para_elem.findall(qn('w:sdt')):
        para_elem.remove(sdt)
    new_run = para_elem.makeelement(qn('w:r'), {})
    if base_rPr is not None:
        new_run.append(deepcopy(base_rPr))
    for i, line in enumerate(new_text.split('\n')):
        if i > 0:
            br = new_run.makeelement(qn('w:br'), {})
            new_run.append(br)
        t = new_run.makeelement(qn('w:t'), {})
        t.text = line
        t.set(qn('xml:space'), 'preserve')
        new_run.append(t)
    para_elem.append(new_run)

def parse_md_table(text):
    lines = [l.strip() for l in text.split('\n') if l.strip()]
    table_lines = [l for l in lines if l.startswith('|') and l.count('|') >= 3]
    if len(table_lines) < 3:
        return None, None, None
    def parse_row(line):
        cells = [c.strip() for c in line.split('|')]
        if cells and cells[0] == '': cells = cells[1:]
        if cells and cells[-1] == '': cells = cells[:-1]
        return cells
    headers = parse_row(table_lines[0])
    rows = []
    for line in table_lines[1:]:
        if re.match(r'^[\|\s\-:]+$', line):
            continue
        rows.append(parse_row(line))
    # Get non-table text
    before = []
    after = []
    in_table = False
    past_table = False
    for line in text.split('\n'):
        stripped = line.strip()
        if stripped.startswith('|') and stripped.count('|') >= 3:
            in_table = True
        elif in_table and not (stripped.startswith('|') and stripped.count('|') >= 3):
            past_table = True
            in_table = False
        if not in_table and not past_table:
            before.append(line)
        elif past_table:
            after.append(line)
    return headers, rows, ('\n'.join(before).strip(), '\n'.join(after).strip())

def insert_word_table(doc, headers, rows, insert_before_elem):
    """Create a Word table using the template's existing table style."""
    ncols = len(headers)
    nrows = len(rows) + 1
    tbl = doc.add_table(rows=nrows, cols=ncols)
    # Use Tabelraster (the template's table style)
    tbl_elem = tbl._tbl
    tblPr = tbl_elem.find(qn('w:tblPr'))
    if tblPr is not None:
        # Remove any default style and set Tabelraster
        for old_style in tblPr.findall(qn('w:tblStyle')):
            tblPr.remove(old_style)
        style_elem = tblPr.makeelement(qn('w:tblStyle'), {})
        style_elem.set(qn('w:val'), 'Tabelraster')
        tblPr.insert(0, style_elem)
    # Fill
    for ci, header in enumerate(headers):
        cell = tbl.rows[0].cells[ci]
        cell.text = header
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(8)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            if ci < ncols:
                cell = tbl.rows[ri+1].cells[ci]
                cell.text = val
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.size = Pt(8)
    # Move table from end of doc to correct position
    body.remove(tbl_elem)
    insert_before_elem.addprevious(tbl_elem)
    return tbl_elem

def fill_section_with_tables(para_elem, section_text):
    """Fill a paragraph, splitting out any markdown tables as Word tables."""
    headers, rows, texts = parse_md_table(section_text)
    if headers is None:
        # No table, just formatted text
        replace_para_formatted(para_elem, section_text)
        return
    before_text, after_text = texts
    # Put before-text in the paragraph
    if before_text:
        replace_para_formatted(para_elem, before_text)
    else:
        replace_para_formatted(para_elem, "")
    # Find next sibling to insert before
    next_elem = para_elem.getnext()
    # If there's after-text, create a paragraph for it
    if after_text:
        new_p = body.makeelement(qn('w:p'), {})
        pPr = para_elem.find(qn('w:pPr'))
        if pPr is not None:
            new_p.append(deepcopy(pPr))
        if next_elem is not None:
            next_elem.addprevious(new_p)
        else:
            body.append(new_p)
        # Format the after-text
        for run in process_formatted_text(after_text.split('\n')[0], new_p, None):
            new_p.append(run)
        # Add remaining lines as line breaks
        for line in after_text.split('\n')[1:]:
            br_run = new_p.makeelement(qn('w:r'), {})
            br = br_run.makeelement(qn('w:br'), {})
            br_run.append(br)
            new_p.append(br_run)
            for run in process_formatted_text(line, new_p, None):
                new_p.append(run)
        insert_word_table(doc, headers, rows, new_p)
    else:
        if next_elem is not None:
            insert_word_table(doc, headers, rows, next_elem)

# ================================================================
# FILL ALL FIELDS
# ================================================================
children = list(body)

# Date
replace_para_plain(children[2], "01/04/2026 (filed), revised 15/04/2026")

# Title
replace_para_formatted(children[4],
    "Geometric Dosimetry: Closed-Form Method, Differentiable Engine, and Interactive Platform for Real-Time Computation and Optimization of Electromagnetic Absorption on Human Bodies\nSoftware platform: AEGIS (Adaptive Electromagnetic Geometric Illumination & Safety)")

# Section 1 questions
fill_section_with_tables(children[9], extract_section(md,
    "### If you would consider your invention as a solution",
    "### How have others tried"))

fill_section_with_tables(children[11], extract_section(md,
    "### How have others tried to solve this problem",
    "### Describe the advantage"))

fill_section_with_tables(children[13], extract_section(md,
    "### Describe the advantage of your solution",
    "### Give a short description"))

fill_section_with_tables(children[15], extract_section(md,
    "### Give a short description of the invention",
    "### If possible, provide a figure"))

replace_para_formatted(children[17], extract_section(md,
    "### If possible, provide a figure",
    "### Does your invention possess disadvantages"))

replace_para_formatted(children[19], extract_section(md,
    "### Does your invention possess disadvantages",
    "### Describe the development status"))

replace_para_formatted(children[21], extract_section(md,
    "### Describe the development status",
    "---\n\n## 2."))

# Disclosure table
table0 = children[24]
t0rows = table0.findall(qn('w:tr'))
answers = {1:"None", 2:"None",
    3:"Planned: JSAC SI 'Digital Twins for Wireless Networks' (deadline May 1, 2026)",
    4:"None",
    5:"Planned: PhD thesis defense before August 2026",
    6:"None", 7:"None", 8:"None", 9:"None"}
for ri in range(1, len(t0rows)):
    cells = t0rows[ri].findall(qn('w:tc'))
    if len(cells) > 1 and ri in answers:
        paras = cells[-1].findall(qn('w:p'))
        if paras:
            replace_para_plain(paras[0], answers[ri])

# Prior art by inventors
replace_para_formatted(children[26], extract_section(md,
    "### Prior art: publications by the inventors",
    "### Prior art: publications by others"))

# Prior art by others
replace_para_formatted(children[28], extract_section(md,
    "### Prior art: publications by others",
    "### Is literature screened"))

# Keywords
replace_para_plain(children[32], extract_section(md,
    "### Keywords:", "### Patents or patent applications"))

# Patents
replace_para_formatted(children[34], extract_section(md,
    "### Patents or patent applications", "### Who are the main"))

# Research groups
replace_para_plain(children[36], extract_section(md,
    "### Who are the main academic or industrial", "---\n\n## 3. Inventors"))

# Inventor 1
replace_para_plain(children[44], "\tRobin\t Wydaeghe")
replace_para_plain(children[45], "\tGhent University / IMEC / INTEC-WAVES")
replace_para_plain(children[46], "\tTechnologiepark-Zwijnaarde 126, 9052 Gent")
replace_para_plain(children[47], "\trobin.wydaeghe@hotmail.com / robin.wydaeghe@ugent.be")
replace_para_plain(children[48], "\t+32 483 06 90 27")
replace_para_plain(children[49], "\tBelgian")
replace_para_plain(children[50], "\t100% -- Sole inventor of the theoretical framework and sole developer of the software platform (AEGIS, ~54,000 lines of code, 2,469 tests)")

# Inventors 2-4: N/A
for inv_start in [55, 66, 77]:
    replace_para_plain(children[inv_start], "\tN/A\t N/A")
    for offset in range(1, 8):
        idx = inv_start + offset
        if idx < len(children):
            tag = children[idx].tag.split('}')[-1] if '}' in children[idx].tag else children[idx].tag
            if tag == 'p':
                replace_para_plain(children[idx], "\tN/A")

# Section 4
replace_para_formatted(children[91], extract_section(md,
    "### Are lab records available?", "### At what site"))
replace_para_plain(children[93], extract_section(md,
    "### At what site", "### Does the invention incorporate"))
replace_para_plain(children[95], extract_section(md,
    "### Does the invention incorporate", "---\n\n## 5."))

# Section 5 funding table
table1 = children[98]
t1rows = table1.findall(qn('w:tr'))
if len(t1rows) > 1:
    cells = t1rows[1].findall(qn('w:tc'))
    for ci, txt in enumerate(["None", "N/A", "N/A"]):
        if ci < len(cells):
            paras = cells[ci].findall(qn('w:p'))
            if paras:
                replace_para_plain(paras[0], txt)

replace_para_formatted(children[100], extract_section(md,
    "### Future funding", "---\n\n## 6."))

# Section 6
replace_para_formatted(children[103], extract_section(md,
    "### Is the invention the result of a collaborative", "### Is there a contract"))
replace_para_plain(children[105], extract_section(md,
    "### Is there a contract", "---\n\n## 7."))

# Section 7
replace_para_formatted(children[108], extract_section(md,
    "### Is there any expression of this invention through software", "### What is the purpose"))
replace_para_formatted(children[110], extract_section(md,
    "### What is the purpose of the software", "### Is the software a derivative"))
replace_para_plain(children[112], extract_section(md,
    "### Is the software a derivative", "---\n\n## 8."))

# Section 8
replace_para_formatted(children[115], extract_section(md,
    "### In your opinion, what kind of commercial applications", "### Which companies could"))
replace_para_plain(children[117], extract_section(md,
    "### Which companies could be interested", "---\n\n## 9."))

# Save - single write
doc.save("/home/user/aegis/spinoff/Invention Disclosure Document.docx")
print("Done! Single-pass fill with tables and formatting.")
