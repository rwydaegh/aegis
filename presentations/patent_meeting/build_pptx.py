#!/usr/bin/env python3
"""Assemble the rendered main.pdf page PNGs into a 16:9 pptx, one PNG per slide."""
import glob
import os

from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
PNG_DIR = os.path.join(HERE, "pdf_png")
OUT = os.path.join(HERE, "main.pptx")

# main.pdf is 16:9 (453.543 x 255.118 pt). Use the standard 13.333 x 7.5 in canvas.
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank = prs.slide_layouts[6]  # fully blank layout

pngs = sorted(glob.glob(os.path.join(PNG_DIR, "slide-*.png")))
for png in pngs:
    slide = prs.slides.add_slide(blank)
    # Each page already matches 16:9, so fill the whole canvas.
    slide.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)

prs.save(OUT)
print(f"Wrote {OUT} with {len(pngs)} slides")
